import tempfile
import unittest
import json
from pathlib import Path

from openpyxl import load_workbook

import database.database as database
from database.database import (
    crear_base_datos,
    guardar_turno_calendario,
    insertar_repartidor,
    insertar_restaurante,
    insertar_turno
)
from services.exportador import (
    crear_calendario_ics,
    exportar_delivery_json,
    exportar_excel,
    exportar_ics
)


class TestExportador(unittest.TestCase):

    def setUp(self):

        self.ruta_original = database.RUTA_BD
        self.temporal = tempfile.TemporaryDirectory()
        database.RUTA_BD = Path(self.temporal.name) / "delivery.db"
        crear_base_datos()

    def tearDown(self):

        database.RUTA_BD = self.ruta_original
        self.temporal.cleanup()

    def test_exportar_excel_no_falla_con_calendario_semanal(self):

        repartidor_id = insertar_repartidor(
            "Ana",
            30,
            "Centro",
            1,
            1,
            50,
            50,
            50,
            descanso_inicio="lunes",
            descanso_fin="martes",
            disponibilidad={"miercoles": "Ambos"}
        )
        restaurante_id = insertar_restaurante(
            "BK Centro",
            "Rua 1",
            "Centro",
            "600000000",
            50
        )
        turno_id = insertar_turno(
            "Comida",
            "Comida",
            "13:00",
            "16:00",
            "#2563EB",
            3
        )
        guardar_turno_calendario(
            "miercoles",
            turno_id,
            restaurante_id,
            repartidor_id,
            "2026-07-13"
        )
        salida = Path(self.temporal.name) / "cuadrante.xlsx"

        exportar_excel(salida, "2026-07-13")

        self.assertTrue(salida.exists())
        self.assertGreater(salida.stat().st_size, 0)

    def test_exportar_excel_incluye_cuadrante_semanal_profesional(self):

        repartidor_id = insertar_repartidor(
            "Ana",
            10,
            "Centro",
            1,
            1,
            50,
            50,
            50,
            descanso_inicio="martes",
            descanso_fin="miercoles",
            disponibilidad={
                dia: "Ambos"
                for dia in (
                    "lunes",
                    "martes",
                    "miercoles",
                    "jueves",
                    "viernes",
                    "sabado",
                    "domingo"
                )
            }
        )
        restaurante_id = insertar_restaurante(
            "Zona Centro",
            "Rua 1",
            "Centro",
            "600000000",
            50
        )
        comida_id = insertar_turno(
            "Comida",
            "Comida",
            "13:00",
            "16:00",
            "#D9F0F2",
            3
        )
        cena_id = insertar_turno(
            "Cena",
            "Cena",
            "20:00",
            "23:30",
            "#D7E2F5",
            3.5
        )
        valle_id = insertar_turno(
            "Comida",
            "Horas valle",
            "00:30",
            "04:30",
            "#DCFCE7",
            4
        )
        guardar_turno_calendario(
            "lunes",
            comida_id,
            restaurante_id,
            repartidor_id,
            "2026-08-10"
        )
        guardar_turno_calendario(
            "lunes",
            cena_id,
            restaurante_id,
            repartidor_id,
            "2026-08-10"
        )
        guardar_turno_calendario(
            "viernes",
            valle_id,
            restaurante_id,
            repartidor_id,
            "2026-08-10"
        )
        salida = Path(self.temporal.name) / "cuadrante_profesional.xlsx"

        exportar_excel(salida, "2026-08-10")

        libro = load_workbook(salida)
        hoja = libro["Cuadrante semanal"]
        self.assertEqual(libro.sheetnames[0], "Cuadrante semanal")
        self.assertIn("Horarios", libro.sheetnames)
        self.assertEqual(hoja["A5"].value, "Ana")
        self.assertEqual(hoja["B5"].value, "10h")
        self.assertEqual(hoja["C4"].value, "Lunes\n10/08")
        self.assertEqual(hoja["D4"].value, "Martes\n11/08")
        self.assertIn("DOBLE", hoja["C5"].value)
        self.assertIn("COMIDA 13:00-16:00 (3 h)", hoja["C5"].value)
        self.assertIn("CENA 20:00-23:30 (3.5 h)", hoja["C5"].value)
        self.assertEqual(hoja["D5"].value, "LIBRE")
        self.assertIn("VALLE", hoja["G5"].value)
        self.assertEqual(hoja["J5"].value, "10.5 h")
        self.assertEqual(hoja["K5"].value, "0.5 h")
        self.assertIn("Leyenda", [celda.value for celda in hoja["A"]])

    def test_exportar_ics_crea_evento_de_calendario(self):

        repartidor_id = insertar_repartidor(
            "Ana",
            30,
            "Centro",
            1,
            1,
            50,
            50,
            50,
            descanso_inicio="lunes",
            descanso_fin="martes",
            disponibilidad={"miercoles": "Ambos"}
        )
        restaurante_id = insertar_restaurante(
            "BK Centro",
            "Rua 1",
            "Centro",
            "600000000",
            50
        )
        turno_id = insertar_turno(
            "Comida",
            "Comida",
            "13:00",
            "16:00",
            "#2563EB",
            3
        )
        guardar_turno_calendario(
            "miercoles",
            turno_id,
            restaurante_id,
            repartidor_id,
            "2026-07-13"
        )
        salida = Path(self.temporal.name) / "cuadrante.ics"

        exportar_ics(salida, "2026-07-13")

        contenido = salida.read_text(encoding="utf-8")
        self.assertIn("BEGIN:VCALENDAR", contenido)
        self.assertIn("BEGIN:VEVENT", contenido)
        self.assertIn("SUMMARY:Comida - BK Centro", contenido)
        self.assertIn("DTSTART:20260715T130000", contenido)
        self.assertIn("DTEND:20260715T160000", contenido)
        self.assertIn("Repartidor: Ana", contenido)

    def test_exportar_delivery_json_crea_payload_generico(self):

        repartidor_id = insertar_repartidor(
            "Ana",
            30,
            "Centro",
            1,
            1,
            50,
            50,
            50,
            descanso_inicio="lunes",
            descanso_fin="martes",
            disponibilidad={"miercoles": "Ambos"}
        )
        restaurante_id = insertar_restaurante(
            "BK Centro",
            "Rua 1",
            "Centro",
            "600000000",
            50
        )
        turno_id = insertar_turno(
            "Comida",
            "Comida",
            "13:00",
            "16:00",
            "#2563EB",
            3
        )
        guardar_turno_calendario(
            "miercoles",
            turno_id,
            restaurante_id,
            repartidor_id,
            "2026-07-13"
        )
        salida = Path(self.temporal.name) / "delivery.json"

        exportar_delivery_json(salida, "2026-07-13")

        payload = json.loads(salida.read_text(encoding="utf-8"))
        self.assertEqual(payload["schema"], "planificador.delivery.v1")
        self.assertEqual(payload["fecha_inicio_semana"], "2026-07-13")
        self.assertEqual(payload["turnos"][0]["restaurante"], "BK Centro")
        self.assertEqual(payload["turnos"][0]["repartidor"], "Ana")

    def test_ics_extiende_turnos_que_cruzan_medianoche(self):

        contenido = crear_calendario_ics({
            "fecha_inicio_semana": "2026-07-13",
            "horarios": [[
                "viernes",
                "Cena",
                "Cena",
                "BK Noche",
                "Centro",
                "Luis",
                "22:00",
                "01:00",
                3
            ]]
        })

        self.assertIn("DTSTART:20260717T220000", contenido)
        self.assertIn("DTEND:20260718T010000", contenido)


if __name__ == "__main__":

    unittest.main()
